import math
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FLAG_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)


def _safe(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)[:32000]
    return v


def _sheet(wb, title, df, flag_col=None):
    ws = wb.create_sheet(title)
    if df.empty:
        ws.append(["No data"])
        return ws
    ws.append(list(df.columns))
    for c in ws[1]:
        c.fill, c.font, c.border = HEADER_FILL, Font(bold=True, color="FFFFFF"), BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r_i, (_, row) in enumerate(df.iterrows(), start=2):
        ws.append([_safe(v) for v in row.tolist()])
        flagged = flag_col and bool(row.get(flag_col))
        for c_i in range(1, len(df.columns) + 1):
            cell = ws.cell(r_i, c_i)
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float)) else "left")
            if flagged:
                cell.fill = FLAG_FILL
    for i, col in enumerate(df.columns, 1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).head(500)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 3, 12), 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def export_report(daily, monthly, customers, customer_summary, products, categories, clean_report,
                   filename="Ecommerce_Analytics_Report.xlsx"):
    output_file = filename if os.path.isabs(filename) or os.path.dirname(filename) else os.path.join(OUTPUT_DIR, filename)
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Summary")
    rows = [
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Orders received", clean_report["rows_in"]),
        ("Rows removed (missing required fields)", clean_report["missing_required_removed"]),
        ("Rows removed (invalid dates)", clean_report["invalid_dates_removed"]),
        ("Rows removed (invalid totals)", clean_report["invalid_totals_removed"]),
        ("Duplicate orders removed", clean_report["duplicate_orders_removed"]),
        ("Clean orders analyzed", clean_report["rows_out"]),
        ("Outliers flagged (unusually high/low order value)", clean_report["outliers_flagged"]),
        ("", ""),
        ("Total customers", customer_summary.get("total_customers", "-")),
        ("Repeat customers", customer_summary.get("repeat_customers", "-")),
        ("One-time customers", customer_summary.get("one_time_customers", "-")),
        ("Average order value", customer_summary.get("avg_order_value_overall", "-")),
    ]
    for label, value in rows:
        ws.append([label, value])
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 22

    _sheet(wb, "Sales_Trends_Daily", daily)
    ws_m = _sheet(wb, "Sales_Trends_Monthly", monthly)
    if not monthly.empty:
        chart = LineChart()
        chart.title = "Monthly Revenue"
        chart.y_axis.title = "Revenue"
        chart.x_axis.title = "Month"
        data = Reference(ws_m, min_col=3, min_row=1, max_row=ws_m.max_row)
        cats = Reference(ws_m, min_col=1, min_row=2, max_row=ws_m.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws_m.add_chart(chart, "F2")

    _sheet(wb, "Customer_Behavior", customers)

    ws_p = _sheet(wb, "Product_Performance", products)
    if not products.empty:
        chart = BarChart()
        chart.title = "Top Products by Revenue"
        chart.y_axis.title = "Revenue"
        data = Reference(ws_p, min_col=2, min_row=1, max_row=ws_p.max_row)
        cats = Reference(ws_p, min_col=1, min_row=2, max_row=ws_p.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws_p.add_chart(chart, "F2")

    if not categories.empty:
        _sheet(wb, "Category_Breakdown", categories)

    try:
        wb.save(output_file)
    except PermissionError:
        base, ext = os.path.splitext(output_file)
        output_file = f"{base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
        wb.save(output_file)
        print("NOTE: The previous file was open in Excel, saved with a new name.")
    return output_file
